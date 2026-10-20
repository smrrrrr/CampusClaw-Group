"""Lesson 2-5 acceptance audit. Run in backend runtime via stdin.

Creates a temporary database/uploads directory before importing the application.
Never seeds, resets passwords, or uploads into the running project's database.
Optional --real-llm performs one short request using the configured server key.
"""
import asyncio
import io
import json
import os
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace

import httpx


async def main():
    with tempfile.TemporaryDirectory(prefix="campusclaw-audit-") as tmp:
        os.environ["DATABASE_URL"] = f"sqlite:///{tmp}/audit.db"
        os.environ["UPLOAD_DIR"] = f"{tmp}/uploads"
        from config import settings
        from database import Base, SessionLocal, engine, run_simple_migrations
        from sqlalchemy import text as sql_text
        from main import app
        from models import AgentConfig, AIQuestionLog, AuditLog, Class, Material, MaterialChunk, Message, Role, User
        from security import hash_password
        from rag_search_service import retrieve_top_k
        import llm_service
        from routers import student

        Base.metadata.create_all(engine)
        password = "AuditOnly@2026"
        with SessionLocal() as db:
            db.add_all([Class(id=i, name=f"Audit class {i}") for i in (101, 102, 103)])
            db.flush()
            for identifier, role, cid, classes, head, force in [
                ("T1", Role.teacher, None, "[101,102]", 101, False),
                ("T2", Role.teacher, None, "[102]", 102, False),
                ("SUBJECT", Role.teacher, None, "[101]", None, False),
                ("S1", Role.student, 101, None, None, False),
                ("S2", Role.student, 102, None, None, False),
                ("FIRST", Role.student, 101, None, None, True),
                ("NOCLASS", Role.student, None, None, None, False),
                ("LOCK", Role.student, 101, None, None, False),
            ]:
                db.add(User(identifier=identifier, role=role, name=identifier,
                            class_id=cid, teaching_classes=classes, subject=("physics" if identifier == "T2" else "math") if classes else None,
                            homeroom_class_id=head, password_hash=hash_password(password),
                            must_change_password=force))
            db.commit()
            ids = {u.identifier: u.id for u in db.query(User)}

        checks = []
        def check(name, ok, detail=""):
            row = {"name": name, "passed": bool(ok), "detail": detail}
            checks.append(row)
            print(json.dumps(row, ensure_ascii=False), flush=True)

        # 在临时库模拟旧版单班配置及无学科会话，验证保留数据且可重复升级。
        with engine.begin() as conn:
            conn.execute(sql_text("DROP TABLE agent_configs"))
            conn.execute(sql_text("CREATE TABLE agent_configs(id INTEGER PRIMARY KEY, class_id INTEGER NOT NULL UNIQUE,teacher_id INTEGER NOT NULL,system_prompt TEXT,tutor_enabled BOOLEAN NOT NULL,updated_at DATETIME NOT NULL)"))
            conn.execute(sql_text("INSERT INTO agent_configs VALUES(1,101,:teacher,'MIGRATION_PROMPT',0,'2026-10-18 00:00:00')"), {"teacher": ids["T1"]})
            conn.execute(sql_text("DROP TABLE conversations"))
            conn.execute(sql_text("CREATE TABLE conversations(id INTEGER PRIMARY KEY,student_id INTEGER NOT NULL,class_id INTEGER NOT NULL,title TEXT NOT NULL,context TEXT,created_at DATETIME NOT NULL,updated_at DATETIME NOT NULL)"))
            conn.execute(sql_text("INSERT INTO conversations VALUES(1,:student,101,'legacy',NULL,'2026-10-18 00:00:00','2026-10-18 00:00:00')"), {"student": ids["S1"]})
        run_simple_migrations()
        run_simple_migrations()
        with SessionLocal() as db:
            from models import Conversation
            migrated = db.get(AgentConfig, 1)
            check("migration.config_preserved", migrated.subject == "math" and migrated.system_prompt == "MIGRATION_PROMPT" and not migrated.tutor_enabled)
            legacy = db.get(Conversation, 1)
            check("migration.legacy_conversation_math", legacy.subject == "math" and legacy.title == "legacy")
            db.delete(migrated); db.delete(legacy); db.commit()

        transport = httpx.ASGITransport(app=app, raise_app_exceptions=False)
        async with httpx.AsyncClient(transport=transport, base_url="http://audit.local") as c:
            async def req(method, url, who=None, **kwargs):
                return await c.request(method, url, headers=headers.get(who, {}), **kwargs)

            headers = {}
            for who in ("T1", "T2", "SUBJECT", "S1", "S2", "FIRST", "NOCLASS"):
                r = await c.post("/api/auth/login", json={"identifier": who, "password": password})
                check(f"auth.login.{who}", r.status_code == 200)
                headers[who] = {"Authorization": f"Bearer {r.json()['access_token']}"}
            check("auth.refresh_cookie_httponly", "HttpOnly" in r.headers.get("set-cookie", ""))
            refresh = await c.post("/api/auth/refresh")
            check("auth.refresh", refresh.status_code == 200)
            check("auth.refresh_restores_full_user", refresh.json().get("user", {}).get("id") == ids["NOCLASS"])
            check("auth.anonymous_denied", (await req("GET", "/api/materials")).status_code == 401)
            check("auth.student_teacher_api_denied", (await req("GET", "/api/teacher/classes", "S1")).status_code == 403)
            check("auth.teacher_student_api_denied", (await req("GET", "/api/student/conversations", "T1")).status_code == 403)
            check("auth.teacher_cross_class_denied", (await req("GET", "/api/teacher/students?class_id=103", "T1")).status_code == 403)
            r = await req("GET", "/api/materials", "FIRST")
            check("auth.first_password_blocks_business", r.status_code == 403, f"HTTP {r.status_code}; expected 403")
            check("auth.first_password_blocks_conversation", (await req("POST", "/api/student/conversations", "FIRST", json={})).status_code == 403)
            check("auth.first_password_blocks_history", (await req("GET", "/api/student/ai-history", "FIRST")).status_code == 403)
            r = await req("POST", "/api/auth/change-password", "FIRST", json={"old_password": password, "new_password": "weak"})
            check("auth.weak_password_denied", r.status_code == 400)
            r = await req("POST", "/api/auth/change-password", "FIRST", json={"old_password": password, "new_password": "ChangedAudit@2026"})
            check("auth.change_password", r.status_code == 200 and not r.json()["must_change_password"])
            for _ in range(5):
                await c.post("/api/auth/login", json={"identifier": "LOCK", "password": "wrong"})
            check("auth.five_failures_lock", (await c.post("/api/auth/login", json={"identifier": "LOCK", "password": password})).status_code == 429)
            check("auth.subject_teacher_reset_denied", (await req("POST", "/api/teacher/reset-password", "SUBJECT", json={"student_id": ids["S1"]})).status_code == 403)
            check("auth.other_homeroom_reset_denied", (await req("POST", "/api/teacher/reset-password", "T2", json={"student_id": ids["S1"]})).status_code == 403)
            check("auth.homeroom_reset", (await req("POST", "/api/teacher/reset-password", "T1", json={"student_id": ids["FIRST"]})).status_code == 200)
            check("auth.reset_immediately_blocks_old_token", (await req("GET", "/api/materials", "FIRST")).status_code == 403)
            with SessionLocal() as db:
                db.add_all([
                    AuditLog(user_id=ids["S1"], role="student", class_id=101, request_path="/audit-class101", method="GET", status_code=403),
                    AuditLog(user_id=ids["S2"], role="student", class_id=102, request_path="/audit-class102", method="GET", status_code=403),
                ])
                db.commit()
            r = await req("GET", "/api/teacher/audit-logs", "T1")
            check("auth.audit_logs_query", r.status_code == 200, f"HTTP {r.status_code}")
            check("auth.audit_logs_class_isolation", r.status_code == 200 and any(x["request_path"] == "/audit-class101" for x in r.json()["items"]) and all(x["request_path"] != "/audit-class102" for x in r.json()["items"]))
            check("auth.audit_subject_teacher_denied", (await req("GET", "/api/teacher/audit-logs", "SUBJECT")).status_code == 403)
            check("auth.audit_other_class_denied", (await req("GET", "/api/teacher/audit-logs?class_id=102", "T1")).status_code == 403)
            csv_text = "学号,学生姓名,班级\nS1,ILLEGAL_CHANGE,Audit class 101\n"
            r = await req("POST", "/api/teacher/import-students", "T2", files={"file": ("students.csv", csv_text.encode(), "text/csv")})
            check("auth.csv_other_class_update_denied", r.status_code == 403, f"HTTP {r.status_code}; teacher only teaches 102, target student belongs to 101")
            csv_text = "学号,学生姓名,班级\nNEW103,New Student,Audit class 103\n"
            r = await req("POST", "/api/teacher/import-students", "T1", files={"file": ("students.csv", csv_text.encode(), "text/csv")})
            check("auth.csv_other_class_create_denied", r.status_code == 403, f"HTTP {r.status_code}; target class 103 is not assigned")
            csv_text = "学号,学生姓名,班级\nS1,ILLEGAL_CLAIM,Audit class 102\n"
            r = await req("POST", "/api/teacher/import-students", "T2", files={"file": ("students.csv", csv_text.encode(), "text/csv")})
            check("auth.csv_cannot_claim_other_student", r.status_code == 403)
            csv_text = "学号,学生姓名\nS1,ILLEGAL_NO_CLASS\n"
            r = await req("POST", "/api/teacher/import-students", "T2", files={"file": ("students.csv", csv_text.encode(), "text/csv")})
            check("auth.csv_missing_class_cannot_bypass", r.status_code == 403)
            csv_text = "学号,学生姓名,班级\nNEW101,New Student,Audit class 101\n"
            r = await req("POST", "/api/teacher/import-students", "T1", files={"file": ("students.csv", csv_text.encode(), "text/csv")})
            check("auth.csv_assigned_class_allowed", r.status_code == 200 and r.json()["created"] == 1)
            csv_text = "学号,学生姓名,班级\nROLLBACK101,New Student,Audit class 101\nNEW103,New Student,Audit class 103\n"
            r = await req("POST", "/api/teacher/import-students", "T1", files={"file": ("students.csv", csv_text.encode(), "text/csv")})
            with SessionLocal() as db:
                check("auth.csv_forbidden_batch_rolled_back", r.status_code == 403 and db.query(User).filter_by(identifier="ROLLBACK101").first() is None)

            text = "# 一元二次方程与求根公式\n一元二次方程的判别式是 $b^2-4ac$。判别式大于零有两个不相等实数根。求根公式用于一元二次方程求解。" * 12
            async def upload(filename, payload, cid=101, who="T1"):
                return await req("POST", "/api/teacher/upload-material", who,
                                 data={"class_id": str(cid), "description": "audit", "tags": "数学"},
                                 files={"file": (filename, payload, "application/octet-stream")})

            r = await upload("quadratic.md", text.encode())
            check("materials.upload_md", r.status_code == 200)
            mid = r.json()["id"]
            r2 = await upload("other-class.md", (text + " OTHER_CLASS_SECRET").encode(), 102)
            othermid = r2.json()["id"]
            physics = await upload("physics.md", (text + " PHYSICS_SOURCE").encode(), 102, "T2")
            physicsmid = physics.json()["id"]
            r = await req("GET", "/api/materials?class_id=102&subject_only=true", "T1")
            check("materials.math_subject_list", r.status_code == 200 and {m["id"] for m in r.json()["items"]} == {othermid})
            r = await req("GET", "/api/materials?class_id=102&subject_only=true", "T2")
            check("materials.physics_subject_list", r.status_code == 200 and {m["id"] for m in r.json()["items"]} == {physicsmid})
            r = await req("GET", "/api/materials?class_id=102", "T1")
            check("materials.all_subject_list_preserved", r.status_code == 200 and {m["id"] for m in r.json()["items"]} == {othermid, physicsmid})
            check("materials.subject_list_cross_class_denied", (await req("GET", "/api/materials?class_id=103&subject_only=true", "T1")).status_code == 403)
            check("materials.student_subject_scope_denied", (await req("GET", "/api/materials?subject_only=true", "S1")).status_code == 403)
            for extension in ("txt", "docx", "pptx", "pdf"):
                buf = io.BytesIO()
                if extension == "docx":
                    from docx import Document
                    doc = Document(); doc.add_paragraph(text); doc.save(buf)
                elif extension == "pptx":
                    from pptx import Presentation
                    pres = Presentation(); slide = pres.slides.add_slide(pres.slide_layouts[1])
                    slide.shapes.title.text = "Quadratic"; slide.placeholders[1].text = text; pres.save(buf)
                elif extension == "pdf":
                    import fitz
                    pdf = fitz.open(); page = pdf.new_page()
                    page.insert_textbox(fitz.Rect(30, 30, 500, 750), "Quadratic equations and discriminant b squared minus four ac. " * 10)
                    buf.write(pdf.tobytes()); pdf.close()
                else:
                    buf.write(text.encode())
                r = await upload(f"sample.{extension}", buf.getvalue())
                with SessionLocal() as db:
                    indexed = r.status_code == 200 and db.get(Material, r.json()["id"]).is_indexed
                check(f"materials.upload_parse_{extension}", indexed, f"HTTP {r.status_code}")
            with SessionLocal() as db:
                chunks = db.query(MaterialChunk).filter_by(material_id=mid).all()
                check("rag.chunks_vectors_persisted", len(chunks) > 1 and all(len(ch.embedding) == 512 * 4 for ch in chunks))
                for mode in ("keyword", "vector", "hybrid"):
                    hits = retrieve_top_k(db, [101], "一元二次方程的判别式和求根公式", mode=mode)
                    check(f"rag.{mode}_class_isolation", bool(hits) and all(h["material_id"] != othermid for h in hits))
                check("rag.irrelevant_no_hits", not retrieve_top_k(db, [101], "食堂营业时间几点钟"))
            check("materials.download", (await req("GET", f"/api/materials/{mid}/download", "S1")).content == text.encode())
            check("materials.preview", (await req("GET", f"/api/materials/{mid}/preview", "S1")).status_code == 200)
            check("materials.cross_class_download_denied", (await req("GET", f"/api/materials/{othermid}/download", "S1")).status_code == 403)
            check("materials.cross_class_list_denied", (await req("GET", "/api/materials?class_id=102", "S1")).status_code == 403)
            check("materials.cross_class_upload_denied", (await upload("denied.md", text.encode(), 103)).status_code == 403)
            check("materials.student_upload_denied", (await upload("denied.md", text.encode(), 101, "S1")).status_code == 403)
            check("materials.forged_pdf_denied", (await upload("fake.pdf", b"not a pdf")).status_code == 400)
            check("materials.reindex", (await req("POST", f"/api/teacher/materials/{mid}/reindex", "T1")).json()["is_indexed"])

            marker = "CUSTOM_PROMPT_AUDIT_MARKER"
            body = {"class_id": 101, "system_prompt": marker, "tutor_enabled": True}
            check("agent.config_save", (await req("PUT", "/api/teacher/agent-config", "T1", json=body)).status_code == 200)
            check("agent.config_read", (await req("GET", "/api/teacher/agent-config?class_id=101", "T1")).json()["system_prompt"] == marker)
            check("agent.cross_class_config_denied", (await req("PUT", "/api/teacher/agent-config", "T2", json=body)).status_code == 403)
            tutor_prompt = llm_service.resolve_system_prompt(SimpleNamespace(system_prompt=marker, tutor_enabled=True))
            check("agent.tutor_overrides_custom_prompt", marker not in tutor_prompt and "苏格拉底" in tutor_prompt)
            check("agent.custom_prompt_without_tutor", marker in llm_service.resolve_system_prompt(SimpleNamespace(system_prompt=marker, tutor_enabled=False)))
            check("agent.unconfigured_uses_default_prompt", llm_service.resolve_system_prompt(None) == llm_service.DEFAULT_SYSTEM_PROMPT, "No saved config uses ordinary default; UI switch becomes effective on save")
            from agent_service import followup_retrieval_query
            topic = "一元二次方程的判别式和求根公式"
            tutor_history = [('user', topic), ('assistant', '这里要先确定哪个系数？[1]'), ('user', '不会')]
            for followup in ['2', 'x=2', '不会', '为什么？', '接下来呢']:
                check(f"agent.followup_query_{followup}", topic in followup_retrieval_query(followup, tutor_history))
            for new_question in ['什么是法向量？', '换一道题：讨论函数单调性', 'What is 2 plus 2?']:
                check(f"agent.new_topic_{new_question}", followup_retrieval_query(new_question, tutor_history) == new_question)
            check("agent.followup_without_anchor_unchanged", followup_retrieval_query('2', [('assistant', topic)]) == '2')
            check("agent.followup_latest_user_topic", followup_retrieval_query('不会', tutor_history + [('user', '什么是法向量？')]).startswith('什么是法向量？'))
            check("agent.long_anchor_not_truncated", followup_retrieval_query('不会', [('user', 'x'*2001)]) == '不会')
            captured_queries = []
            original_retrieve = student._retrieve_context
            def observed_retrieve(db, cid, query, subject):
                captured_queries.append((cid, query, subject))
                return original_retrieve(db, cid, query, subject)
            student._retrieve_context = observed_retrieve
            convo = (await req("POST", "/api/student/conversations", "S1", json={"title": "audit"})).json()["id"]
            check("agent.other_student_conversation_hidden", (await req("GET", f"/api/student/conversations/{convo}/messages", "S2")).status_code in (403, 404))
            check("agent.unassigned_student_denied", (await req("POST", "/api/student/conversations", "NOCLASS", json={})).status_code == 403)
            captured = []
            original_stream = student.stream_chat
            async def fake_stream(messages):
                captured.append(messages)
                yield "先判断判别式"
                yield "，你认为应该怎么计算？[1]"
            student.stream_chat = fake_stream
            def events(r):
                return [json.loads(line[6:]) for line in r.text.splitlines() if line.startswith("data: ")]
            r = await req("POST", f"/api/student/conversations/{convo}/chat", "S1", json={"question": "一元二次方程的判别式和求根公式"})
            ev = events(r)
            check("agent.sse_transport_mock_llm", r.headers.get("content-type", "").startswith("text/event-stream") and [e["type"] for e in ev] == ["delta", "delta", "done"])
            check("agent.citations_provenance", bool(ev[-1]["citations"]) and all(x["material_id"] != othermid and x["excerpt"] for x in ev[-1]["citations"]))
            from citation_service import validate_references
            with SessionLocal() as db:
                _, candidates = student._retrieve_context(db, 101, "一元二次方程的判别式和求根公式", "math")
                source = ev[-1]["citations"][0]
                original = db.query(MaterialChunk).filter_by(material_id=source["material_id"], chunk_index=source["chunk_index"]).one()
                check("citations.excerpt_is_exact_chunk", source["excerpt"] == original.content)
            check("citations.only_used_source_attached", len(ev[-1]["citations"]) == 1 and ev[-1]["citations"][0]["reference_id"] == 1)
            chosen = candidates[-1]
            checked, used = validate_references(f"核对原文[{chosen.reference_id}]", candidates)
            check("citations.original_number_preserved", len(used) == 1 and used[0].reference_id == chosen.reference_id and f"[{chosen.reference_id}]" in checked)
            checked, used = validate_references("伪造出处[999]", candidates)
            check("citations.invalid_number_removed", "[999]" not in checked and not used and "尚未核实" in checked)
            checked, used = validate_references("无资料仍声称来源[1]", [])
            check("citations.no_context_cannot_cite_history", "[1]" not in checked and not used)
            checked, used = validate_references("普通引导，不使用引用。", candidates)
            check("citations.uncited_candidates_not_attached", not used)
            checked, used = validate_references("来源[1](https://fabricated.example/source)", candidates)
            check("citations.model_url_removed", "fabricated.example" not in checked and used[0].reference_id == 1)
            prompt = llm_service.build_chat_messages("CUSTOM", [("assistant", "旧编号[1]")], "新问题", "[1]本轮原文")
            check("citations.policy_applies_to_custom", llm_service.CITATION_POLICY in prompt[0]["content"] and "[1]" not in prompt[1]["content"])
            check("citations.original_version_download", (await req("GET", f"/api/materials/{source['material_id']}/download?expected_sha256={source['file_sha256']}", "S1")).status_code == 200)
            check("citations.wrong_version_rejected", (await req("GET", f"/api/materials/{source['material_id']}/download?expected_sha256={'0'*64}", "S1")).status_code == 409)
            check("agent.actual_messages_use_tutor_override", marker not in captured[0][0]["content"] and "苏格拉底" in captured[0][0]["content"])
            await req("POST", f"/api/student/conversations/{convo}/chat", "S1", json={"question": "接下来呢"})
            check("agent.multiturn_context", len(captured[-1]) == 4 and "判别式" in captured[-1][2]["content"])
            check("agent.followup_retrieves_original_topic_with_same_scope", captured_queries[-1] == (101, topic+'\n学生追问：接下来呢', 'math'))
            check("agent.followup_preserves_raw_student_message", captured[-1][-1]['content'].startswith('问题：接下来呢\n'))
            r = await req("POST", f"/api/student/conversations/{convo}/chat", "S1", json={"question": "2"})
            check("agent.numeric_followup_refetches_sources", r.status_code == 200 and captured_queries[-1] == (101, topic+'\n学生追问：2', 'math') and bool(events(r)[-1]['citations']))
            student._retrieve_context = original_retrieve
            with SessionLocal() as db:
                check("agent.messages_persist", db.query(Message).filter_by(conversation_id=convo).count() == 6)
            async def broken_stream(messages):
                raise llm_service.LLMGenerationError("audit simulated failure")
                yield ""
            student.stream_chat = broken_stream
            r = await req("POST", f"/api/student/conversations/{convo}/chat", "S1", json={"question": "一元二次方程的判别式"})
            ev = events(r)
            check("agent.failure_fallback", any(e["type"] == "error" and "相关段落" in e["message"] for e in ev))
            history = (await req("GET", f"/api/student/conversations/{convo}/messages", "S1")).json()
            check("agent.failure_sources_persisted", bool(history[-1]["sources"]), "Error event has text only; structured sources are available in persisted history")
            student.stream_chat = original_stream
            with SessionLocal() as db:
                count = db.query(AIQuestionLog).filter_by(student_id=ids["S1"]).count()
                db.add_all(AIQuestionLog(student_id=ids["S1"], class_id=101, question="audit quota") for _ in range(40-count))
                db.commit()
            student.stream_chat = fake_stream
            r = await req("POST", f"/api/student/conversations/{convo}/chat", "S1", json={"question": "一元二次方程的判别式"})
            check("agent.spec_40_question_limit", r.status_code == 429, f"HTTP {r.status_code}; 40 per China calendar day")
            student.stream_chat = original_stream
            # 同班数学与物理教师分别配置，学生按科目切换配置/历史/资料。
            for who, subject, marker in [("T1", "math", "MATH_OWN_PROMPT"), ("T2", "physics", "PHYSICS_OWN_PROMPT")]:
                saved = await req("PUT", "/api/teacher/agent-config", who, json={"class_id": 102, "system_prompt": marker, "tutor_enabled": False})
                check(f"agent.subject_save_{subject}", saved.status_code == 200 and saved.json()["subject"] == subject)
            read_math = (await req("GET", "/api/teacher/agent-config?class_id=102", "T1")).json()
            check("agent.physics_save_does_not_overwrite_math", read_math["system_prompt"] == "MATH_OWN_PROMPT")
            with SessionLocal() as db:
                db.get(User, ids["T2"]).subject = "english"; db.commit()
            changed = (await req("GET", "/api/teacher/agent-config?class_id=102", "T2")).json()
            check("agent.teacher_subject_uses_current_database", changed["subject"] == "english" and changed["system_prompt"] == "")
            with SessionLocal() as db:
                db.get(User, ids["T2"]).subject = "physics"; db.commit()
            own = await req("PUT", "/api/teacher/agent-config", "SUBJECT", json={"class_id": 101, "system_prompt": "OTHER_MATH_TEACHER", "tutor_enabled": False})
            read_original = (await req("GET", "/api/teacher/agent-config?class_id=101", "T1")).json()
            check("agent.same_subject_teacher_configs_independent", own.status_code == 200 and read_original["system_prompt"] == "CUSTOM_PROMPT_AUDIT_MARKER")
            student.stream_chat = fake_stream
            subject_convos = {}
            for subject, teacher, prompt, material_id in [("math", "T1", "MATH_OWN_PROMPT", othermid), ("physics", "T2", "PHYSICS_OWN_PROMPT", physicsmid)]:
                config = (await req("GET", f"/api/student/assistant-config?subject={subject}", "S2")).json()
                check(f"agent.student_config_{subject}", config["configured"] and config["teacher_id"] == ids[teacher] and not config["tutor_enabled"])
                cid = (await req("POST", "/api/student/conversations", "S2", json={"subject": subject})).json()["id"]
                subject_convos[subject] = cid
                response = await req("POST", f"/api/student/conversations/{cid}/chat", "S2", json={"question": "一元二次方程的判别式和求根公式"})
                done = events(response)[-1]
                check(f"agent.prompt_applied_{subject}", prompt in captured[-1][0]["content"])
                check(f"agent.rag_subject_isolation_{subject}", bool(done["citations"]) and all(x["material_id"] == material_id for x in done["citations"]))
            for subject, cid in subject_convos.items():
                listing = (await req("GET", f"/api/student/conversations?subject={subject}", "S2")).json()
                check(f"agent.history_isolated_{subject}", [x["id"] for x in listing] == [cid] and all(x["subject"] == subject for x in listing))
            check("agent.config_cannot_mix_conversation_subject", (await req("GET", f"/api/student/assistant-config?subject=physics&conversation_id={subject_convos['math']}", "S2")).status_code == 404)
            fallback = (await req("GET", "/api/student/assistant-config?subject=english", "S2")).json()
            check("agent.unconfigured_subject_reports_default", not fallback["configured"] and fallback["teacher_id"] is None and not fallback["tutor_enabled"])
            check("agent.invalid_subject_rejected", (await req("POST", "/api/student/conversations", "S2", json={"subject": "unknown"})).status_code == 422)
            student.stream_chat = original_stream
            stale = await upload("光电效应.md", "光电效应的截止频率由材料逸出功决定。光电效应截止频率与入射光频率的关系需要核对原文。".encode())
            stale_id = stale.json()["id"]
            from material_service import resolve_abs_path
            with SessionLocal() as db:
                material = db.get(Material, stale_id)
                path = resolve_abs_path(material.file_path)
                original_bytes = path.read_bytes()
                source_hash = material.source_sha256
                path.write_bytes(original_bytes + b" changed")
                _, refs = student._retrieve_context(db, 101, "光电效应的截止频率", "math")
                check("citations.changed_file_stale_index_excluded", not refs)
            check("citations.changed_original_download_rejected", (await req("GET", f"/api/materials/{stale_id}/download?expected_sha256={source_hash}", "S1")).status_code == 409)
            path.write_bytes(original_bytes)
            with SessionLocal() as db:
                db.get(Material, stale_id).source_sha256 = None; db.commit()
                _, refs = student._retrieve_context(db, 101, "光电效应的截止频率", "math")
                check("citations.unversioned_index_excluded", not refs)
            await req("POST", f"/api/teacher/materials/{stale_id}/reindex", "T1")
            with SessionLocal() as db:
                _, refs = student._retrieve_context(db, 101, "光电效应的截止频率", "math")
                check("citations.reindex_restores_verified_sources", bool(refs) and refs[0].material_id == stale_id)
            check("materials.delete", (await req("DELETE", f"/api/teacher/materials/{mid}", "T1")).status_code == 200)
            with SessionLocal() as db:
                check("materials.delete_cleans_chunks", db.query(MaterialChunk).filter_by(material_id=mid).count() == 0)
            check("auth.logout_clears_cookie", "Max-Age=0" in (await req("POST", "/api/auth/logout", "S1")).headers.get("set-cookie", ""))

        if "--real-llm" in sys.argv:
            if settings.deepseek_api_key:
                try:
                    chunks = []
                    async def collect():
                        async for part in llm_service.stream_chat([
                            {"role": "system", "content": "你是数学学习助手。只问一个引导性问题，使用[1]标记来源，不要直接给答案，最多60字。"},
                            {"role": "user", "content": "[1]一元二次方程判别式为 b²-4ac。问题：判别式如何计算？"},
                        ], timeout=20):
                            chunks.append(part)
                    await asyncio.wait_for(collect(), timeout=40)
                    check("llm.real_deepseek_stream", len(chunks) > 1 and bool("".join(chunks).strip()), f"{len(chunks)} deltas; answer={''.join(chunks)[:120]}")
                except Exception as exc:
                    check("llm.real_deepseek_stream", False, type(exc).__name__)
            else:
                check("llm.real_deepseek_stream", False, "Server API key absent")
        print(json.dumps({"total": len(checks), "passed": sum(x["passed"] for x in checks), "failed": [x for x in checks if not x["passed"]]}, ensure_ascii=False), flush=True)
        engine.dispose()
        return 1 if any(not x["passed"] for x in checks) else 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
